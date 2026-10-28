#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""verify.py — Office 交付件的「结构体检」。

它不能替代肉眼，但能挡住 90% 的低级翻车：空文件、只有标题没有正文、
表格只有表头、公式没有缓存值、PPT 有 8 页但 3 页是空的、16:9 设成了 4:3、
字体没设东亚字体（Office 打开变 Times New Roman）、代码块里的页脚页码其实是死字。

用法：
  python verify.py <file> [--json]
  python verify.py <file> --render-dir <render输出目录>   # 顺便看渲染出来的 PNG 有没有空白页
退出码：0 = 没问题（可能有警告），1 = 有 error。
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Pt

ERR, WARN = "error", "warning"


def _rep(path: str) -> dict:
    return {"file": os.path.abspath(path), "sizeKB": round(os.path.getsize(path) / 1024, 1),
            "issues": [], "info": {}}


def _add(r: dict, level: str, msg: str, where: str = "") -> None:
    r["issues"].append({"level": level, "where": where, "msg": msg})


# ── docx ─────────────────────────────────────────────────────────────────
def verify_docx(path: str) -> dict:
    r = _rep(path)
    d = Document(path)
    paras = d.paragraphs
    text_chars = sum(len(p.text.strip()) for p in paras)
    heads = [p for p in paras if (p.style.name or "").startswith("Heading")]
    empties = sum(1 for p in paras if not p.text.strip())
    tables = d.tables
    imgs = len(d.inline_shapes)
    sec = d.sections[0]
    r["info"].update({
        "paragraphs": len(paras), "textChars": text_chars, "headings": len(heads),
        "headingLevels": sorted({p.style.name for p in heads}),
        "tables": len(tables), "tableCells": sum(len(t.rows) * len(t.columns) for t in tables),
        "images": imgs, "sections": len(d.sections),
        "pageSizeCm": [round(sec.page_width.cm, 1), round(sec.page_height.cm, 1)],
        "marginCm": round(sec.left_margin.cm, 2),
    })
    if text_chars < 50:
        _add(r, ERR, f"正文只有 {text_chars} 字，基本是空文档")
    if not heads and len(paras) > 3:
        _add(r, WARN, "没有任何标题（Heading），导航窗格/目录会空")
    if empties > max(6, len(paras) // 3):
        _add(r, WARN, f"空段落 {empties}/{len(paras)}，排版可能是靠回车堆出来的")
    for i, t in enumerate(tables):
        if len(t.rows) < 2:
            _add(r, ERR, f"第 {i+1} 个表格只有表头没有数据行")
        if not (t.style and t.style.name and "Grid" in t.style.name or t.style):
            _add(r, WARN, f"第 {i+1} 个表格没有套样式，Office 里会是裸线")
    # 东亚字体：没设 w:eastAsia 的中文 run 在 Word 里会回退
    runs = [run for p in paras for run in p.runs]
    cn_runs = [run for run in runs if any("\u4e00" <= ch <= "\u9fff" for ch in run.text)]
    missing_ea = 0
    for run in cn_runs:
        rpr = run._element.find(qn("w:rPr"))
        rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rfonts is None or not rfonts.get(qn("w:eastAsia")):
            missing_ea += 1
    r["info"]["cnRuns"] = len(cn_runs)
    if cn_runs and missing_ea:
        _add(r, WARN, f"{missing_ea}/{len(cn_runs)} 段中文 run 没设东亚字体（w:eastAsia），Word 可能回退成 Times New Roman")
    # 页脚页码域
    footer_xml = ""
    for s in d.sections:
        for p in s.footer.paragraphs:
            footer_xml += p._p.xml
    if "PAGE" in footer_xml:
        r["info"]["footerPageField"] = True
    else:
        _add(r, WARN, "页脚没有 PAGE 页码域（写死数字的话增删页就对不上）")
    return r


# ── xlsx ─────────────────────────────────────────────────────────────────
def verify_xlsx(path: str) -> dict:
    r = _rep(path)
    wb = load_workbook(path, data_only=False)
    wb_vals = load_workbook(path, data_only=True)
    sheets = []
    total_formulas = 0
    missing_cache = 0
    for ws in wb.worksheets:
        fcells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        total_formulas += len(fcells)
        vals_ws = wb_vals[ws.title]
        for c in fcells:
            if vals_ws[c.coordinate].value is None:
                missing_cache += 1
        sheets.append({
            "name": ws.title, "dims": ws.dimensions, "rows": ws.max_row, "cols": ws.max_column,
            "formulas": len(fcells), "charts": len(ws._charts), "merged": len(ws.merged_cells.ranges),
            "freeze": str(ws.freeze_panes or ""), "condFormats": len(list(ws.conditional_formatting)),
            "images": len(getattr(ws, "_images", [])),
        })
        if ws.max_row <= 1:
            _add(r, ERR, f"工作表「{ws.title}」只有表头没有数据")
    r["info"].update({"sheets": sheets, "formulas": total_formulas, "formulaCellsWithoutCache": missing_cache})
    if missing_cache:
        _add(r, WARN, f"{missing_cache}/{total_formulas} 个公式单元格没有缓存值（openpyxl 只写公式不算值）——"
                      "用 Excel/LibreOffice 打开会自动重算，但发给人之前最好打开一次确认数对")
    if not any(s["charts"] for s in sheets):
        _add(r, WARN, "没有任何图表（如果需求里有'可视化'就不合格）")
    charts = sum(s["charts"] for s in sheets)
    r["info"]["charts"] = charts
    return r


# ── pptx ─────────────────────────────────────────────────────────────────
def verify_pptx(path: str) -> dict:
    r = _rep(path)
    prs = Presentation(path)
    w, h = prs.slide_width, prs.slide_height
    ratio = w / h
    slides = []
    empty = []
    small_fonts = 0
    overflow = []
    for i, slide in enumerate(prs.slides, start=1):
        texts, pics, tables, charts = [], 0, 0, 0
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
                for p in sh.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size and run.font.size < Pt(12):
                            small_fonts += 1
                # 粗估溢出：字符数 × 字号 → 需要的行数 vs 可用高度
                for p in sh.text_frame.paragraphs:
                    txt = "".join(run.text for run in p.runs)
                    if not txt:
                        continue
                    size = (p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else 18)
                    box_w_in = sh.width / 914400
                    per_line = max(1, int(box_w_in * 72 / size * 1.7))
                    lines = max(1, -(-len(txt) // per_line))
                    need_in = lines * size * 1.35 / 72
                    if need_in > (sh.height / 914400) * 1.02:
                        overflow.append({"slide": i, "needIn": round(need_in, 2),
                                         "boxIn": round(sh.height / 914400, 2), "text": txt[:40]})
            if sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
                pics += 1
            if getattr(sh, "has_table", False) and sh.has_table:
                tables += 1
            if getattr(sh, "has_chart", False) and sh.has_chart:
                charts += 1
        if not texts:
            empty.append(i)
        slides.append({"page": i, "shapes": len(slide.shapes), "chars": sum(len(t) for t in texts),
                       "pictures": pics, "tables": tables, "charts": charts})
    r["info"].update({
        "slides": len(prs.slides), "sizeIn": [round(w / 914400, 3), round(h / 914400, 3)],
        "ratio": round(ratio, 4), "is16x9": abs(ratio - 16 / 9) < 0.01,
        "perSlide": slides, "totals": {"pictures": sum(s["pictures"] for s in slides),
                                        "tables": sum(s["tables"] for s in slides),
                                        "charts": sum(s["charts"] for s in slides)},
    })
    if len(prs.slides) == 0:
        _add(r, ERR, "0 页")
    if empty:
        _add(r, ERR, f"这些页一个字都没有：{empty}")
    if abs(ratio - 16 / 9) > 0.01:
        _add(r, WARN, f"页面比例 {ratio:.3f} 不是 16:9（13.333×7.5in），投影/手机端会留黑边")
    if overflow:
        r["info"]["overflowSuspects"] = overflow[:8]
        _add(r, WARN, f"{len(overflow)} 处文本框疑似装不下文字（粗估，看渲染图确认）")
    if small_fonts:
        _add(r, WARN, f"{small_fonts} 处字号小于 12pt，投影时看不清")
    if r["info"]["totals"]["charts"] == 0 and len(prs.slides) >= 3:
        _add(r, WARN, "没有任何图表（≥3 页的汇报通常至少一张）")
    return r


# ── 渲染结果 ─────────────────────────────────────────────────────────────
def verify_render(render_dir: str) -> list[dict]:
    issues = []
    rep = os.path.join(render_dir, "render-report.json")
    if not os.path.exists(rep):
        issues.append({"level": WARN, "where": render_dir, "msg": "没有 render-report.json（没渲染过）"})
        return issues
    with open(rep, encoding="utf-8") as f:
        data = json.load(f)
    for it in data.get("items", []):
        if it.get("blank"):
            issues.append({"level": ERR, "where": f"第 {it['page']} 页", "msg": "渲染出来是空白页"})
        elif it.get("inkRatio", 1) < 0.01:
            issues.append({"level": WARN, "where": f"第 {it['page']} 页",
                           "msg": f"ink 只有 {it['inkRatio']:.3f}，几乎是空的"})
        if it.get("edgeInkRatio", 0) > 0.06:
            issues.append({"level": WARN, "where": f"第 {it['page']} 页",
                           "msg": "内容贴近/超出页边（可能跑版）"})
        if it.get("textChars") == 0 and it.get("images") == 0:
            issues.append({"level": WARN, "where": f"第 {it['page']} 页", "msg": "既没有文字也没有图"})
    return issues


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--render-dir", default="")
    a = ap.parse_args()
    ext = os.path.splitext(a.file)[1].lower()
    fn = {".docx": verify_docx, ".doc": verify_docx, ".xlsx": verify_xlsx, ".pptx": verify_pptx}.get(ext)
    if not fn:
        print(f"不支持的类型：{ext}", file=sys.stderr)
        sys.exit(2)
    r = fn(a.file)
    if a.render_dir:
        r["issues"] += verify_render(a.render_dir)
    errors = [i for i in r["issues"] if i["level"] == ERR]
    warns = [i for i in r["issues"] if i["level"] == WARN]
    r["errors"] = len(errors)
    r["warnings"] = len(warns)
    if a.json:
        print(json.dumps(r, ensure_ascii=False, indent=2))
    else:
        print(f"=== {os.path.basename(a.file)}  {r['sizeKB']}KB ===")
        print(json.dumps(r["info"], ensure_ascii=False, indent=2))
        for i in r["issues"]:
            mark = "✗" if i["level"] == ERR else "!"
            print(f"  {mark} {i['where']}: {i['msg']}")
        print(f"结论：{len(errors)} error / {len(warns)} warning")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
