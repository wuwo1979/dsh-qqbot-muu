#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_report.py — 一份数据 → docx 报告 + pptx 汇报 + xlsx 底稿（三件套一次出）。

设计要点（改之前先读 SKILL.md）：
1. **先算再写**：每份数据先过 ec_data 判体系、算指标、收集 warnings/questions，
   报告里的每个结论都必须能指回某个指标，不许写"性能优异"这种没有数字的话。
2. **口径跟着数字走**：每个体系末尾都写"口径与假设"，让人知道这个数是按什么算的。
3. 图由 plot.py 现画（本机没有 matplotlib），Word/PPT 嵌 PNG。

用法：
  python make_report.py --title "LFP 半电池电化学性能" --in <数据目录> --out <输出目录>
         [--mass 2.13] [--area 1.13] [--author 小蓝] [--no-deck]
"""
from __future__ import annotations

import argparse
import glob
import os
import sys

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import ec_data  # noqa: E402
import plot as P  # noqa: E402
from office_kit import Deck, Doc, Xlsx  # noqa: E402

KIND_ORDER = {"cycle": 0, "rate": 1, "gcd": 2, "cv": 3, "eis": 4, "unknown": 9}
KIND_TITLE = {
    "cycle": "循环性能", "rate": "倍率性能", "gcd": "恒流充放电", "cv": "循环伏安（CV）",
    "eis": "交流阻抗（EIS）", "unknown": "未识别数据",
}


def figure_for(r, df: pd.DataFrame, out_png: str) -> str | None:
    """按体系挑最该看的那张图。"""
    m = r.mapping
    try:
        if r.kind == "cycle" and m.get("cycle") and m.get("capacity"):
            t = r.tables.get("每圈容量")
            xs = list(t["cycle"]) if t is not None and len(t) else list(df[m["cycle"]])
            ys = list(t["max"]) if t is not None and len(t) else list(df[m["capacity"]])
            series = {"比容量": (xs, ys)}
            y2 = None
            if m.get("efficiency"):
                ce = [float(v) for v in pd.to_numeric(df[m["efficiency"]], errors="coerce")]
                series["库伦效率%"] = (list(df[m["cycle"]]), ce)
                y2 = "库伦效率%"
            return P.line(series, f"{KIND_TITLE['cycle']}（{len(xs)} 圈）", "循环序号",
                          f"比容量（{r.units.get('capacity') or 'mAh/g'}）", out=out_png, y2=y2)
        if r.kind == "rate":
            t = r.tables.get("倍率容量")
            if t is not None and len(t):
                return P.bars(list(t["rate"]), list(t["q"]), KIND_TITLE["rate"], "倍率",
                              f"比容量（{r.units.get('capacity') or 'mAh/g'}）", out=out_png)
        if r.kind == "eis" and m.get("zreal"):
            x = pd.to_numeric(df[m["zreal"]], errors="coerce")
            y = pd.to_numeric(df[m["zimag"]], errors="coerce")
            ok = x.notna() & y.notna()
            return P.scatter(list(x[ok]), list(y[ok]), "EIS Nyquist", "Z'(Ω)", "-Z''(Ω)", out=out_png)
        if r.kind == "gcd" and m.get("potential") and m.get("time"):
            t = pd.to_numeric(df[m["time"]], errors="coerce")
            u = pd.to_numeric(df[m["potential"]], errors="coerce")
            ok = t.notna() & u.notna()
            return P.line({"电位": (list(t[ok]), list(u[ok]))}, KIND_TITLE["gcd"], "时间(s)", "电位(V)",
                          out=out_png)
        if r.kind == "cv" and m.get("potential") and m.get("current"):
            e = pd.to_numeric(df[m["potential"]], errors="coerce")
            i = pd.to_numeric(df[m["current"]], errors="coerce")
            ok = e.notna() & i.notna()
            return P.line({"电流": (list(e[ok]), list(i[ok]))}, KIND_TITLE["cv"], "电位(V)", "电流(mA)",
                          out=out_png)
    except Exception as e:  # noqa: BLE001  出图失败不能把报告搞崩
        print(f"  ! 出图失败（{r.kind}）：{e}")
    return None


def headline(r) -> str:
    """一句话结论：只说数据里真有的事。"""
    m = r.metrics
    if r.kind == "cycle":
        bits = []
        if "首圈容量" in m:
            bits.append(f"首圈 {m['首圈容量']:g}")
        if "末圈容量" in m:
            bits.append(f"末圈 {m['末圈容量']:g}")
        if "容量保持率%" in m:
            bits.append(f"保持率 {m['容量保持率%']:g}%")
        if "平均库伦效率%" in m:
            bits.append(f"平均库伦效率 {m['平均库伦效率%']:g}%")
        return "；".join(bits)
    if r.kind == "rate":
        return (f"最高倍率容量 {m.get('最高倍率容量', '—')}，"
                f"相对首档保持 {m.get('相对首档保持%', '—')}%")
    if r.kind == "eis":
        return f"Rs≈{m.get('Rs(高频截距,Ω)', '—')} Ω，半圆直径≈{m.get('半圆直径≈Rct(Ω)', '—')} Ω"
    if r.kind == "gcd":
        return (f"电位 {m.get('电位范围V', '—')} V，平均 {m.get('平均电压V', '—')} V，"
                f"容量跨度 {m.get('容量跨度', '—')}")
    if r.kind == "cv":
        return f"ΔEp={m.get('峰电位差ΔEp(V)', '—')} V，|Ipa/Ipc|={m.get('峰电流比|Ipa/Ipc|', '—')}"
    return "未能识别，见下方口径与问题"


def build_docx(items, title: str, out_path: str, author: str = "小蓝") -> str:
    doc = Doc(title, subtitle=f"数据自动解读 · {len(items)} 个数据集", footer=f"{author} · 自动生成")
    doc.heading("一、结论摘要", 1)
    doc.para("下列结论全部来自本次数据的自动计算；口径（单位、取数方式）逐条列在每节末尾，"
             "没确认到的地方在「待确认」里列出，不会替用户猜。")
    rows = []
    for r, _df, _fig in items:
        rows.append([KIND_TITLE.get(r.kind, r.kind), os.path.basename(getattr(r, "source", "")),
                     headline(r) or "—"])
    doc.table(["数据集", "文件", "关键结论"], rows, caption="表 1 结论摘要", widths_cm=[3.0, 5.0, 7.4])

    for idx, (r, _df, fig) in enumerate(items, start=2):
        doc.heading(f"{idx - 1}、{KIND_TITLE.get(r.kind, r.kind)}", 1)
        if fig and os.path.exists(fig):
            doc.image(fig, width_cm=15.4, caption=f"图 {idx - 1} {KIND_TITLE.get(r.kind, '')}")
        if r.metrics:
            doc.table(["指标", "值"],
                      [[k, (f"{v:g}" if isinstance(v, (int, float)) else str(v))] for k, v in r.metrics.items()],
                      caption=f"表 {idx} {KIND_TITLE.get(r.kind, '')}指标", widths_cm=[6.4, 6.4])
        if r.assumptions:
            doc.heading("口径与假设", 2)
            doc.bullets(r.assumptions)
        if r.warnings or r.questions:
            doc.heading("待确认 / 需要补的信息", 2)
            doc.bullets([f"⚠ {w}" for w in r.warnings] + [f"❓ {q}" for q in r.questions])
        if r.unknown:
            doc.para(f"（未识别的列：{'、'.join(str(c) for c in r.unknown)}）", indent=False, size=9)

    doc.heading("附录：方法", 1)
    doc.bullets([
        "体系判定顺序：EIS（频率+实/虚部）→ 循环（循环+容量）→ 倍率（倍率+容量）→ "
        "恒流充放电（时间+电位）→ CV（电位+电流）。",
        "循环数据每圈取该圈容量的最大值，衰减用线性拟合斜率，保持率=末圈/首圈。",
        "能量按 ∫U·dQ 梯形积分；EIS 的 Rs 取最高频实部，Rct 取 max(Z')−Rs（精确值需等效电路拟合）。",
        "所有数值都能追溯回原始文件，清洗后的表在配套 xlsx 的「清洗后数据」页。",
    ])
    return doc.save(out_path)


def build_deck(items, title: str, out_path: str, author: str = "小蓝") -> str:
    deck = Deck(title, subtitle=f"电化学数据汇报 · {len(items)} 组测试 · {author}")
    deck.bullets("一页看完", [f"{KIND_TITLE.get(r.kind, r.kind)}：{headline(r)}" for r, _d, _f in items])
    # 先放一张"可编辑的原生图表"（PNG 糊了也能在 PPT 里改数据），再放各体系的实拍图
    for r, df, _fig in items:
        if r.kind == "rate" and r.tables.get("倍率容量") is not None:
            t = r.tables["倍率容量"]
            deck.chart("倍率性能（可在 PPT 里改数）", [str(x) for x in t["rate"]],
                       {"比容量": [float(v) for v in t["q"]]}, kind="bar", unit="mAh/g")
            break
    for r, df, _fig in items:
        if r.kind == "cycle" and r.tables.get("每圈容量") is not None:
            t = r.tables["每圈容量"]
            step = max(1, len(t) // 24)
            sub = t.iloc[::step]
            deck.chart("循环容量衰减", [str(int(x)) for x in sub["cycle"]],
                       {"比容量": [float(v) for v in sub["max"]]}, kind="line", unit="mAh/g")
            break
    for r, _df, fig in items:
        if fig and os.path.exists(fig):
            deck.picture(KIND_TITLE.get(r.kind, r.kind), fig, caption=headline(r))
    rows = [[KIND_TITLE.get(r.kind, r.kind), os.path.basename(getattr(r, "source", "")), headline(r) or "—"]
            for r, _d, _f in items]
    deck.table("指标一览", ["数据集", "文件", "关键结论"], rows, col_widths=[3.2, 4.6, 3.9])
    return deck.save(out_path)


def build_xlsx(items, out_path: str) -> str:
    from openpyxl.utils import get_column_letter

    xl = Xlsx()
    summary = []
    for r, _df, _fig in items:
        for k, v in r.metrics.items():
            summary.append([KIND_TITLE.get(r.kind, r.kind), os.path.basename(getattr(r, "source", "")), k, v])
    xl.sheet("指标汇总", ["数据集", "文件", "指标", "值"], summary,
             note="所有指标均由脚本按页面口径算出；原始值见各数据集页")

    for r, df, _fig in items:
        clean = r.tables.get("清洗后数据", df)
        head = [str(c) for c in clean.columns]
        name = os.path.splitext(os.path.basename(getattr(r, "source", "data")))[0][:24]
        ws = xl.sheet(name, head, clean.head(3000).itertuples(index=False, name=None),
                     note=f"口径：{KIND_TITLE.get(r.kind, r.kind)}" +
                          ("；" + "；".join(r.notes) if r.notes else ""))
        anchor = f"{get_column_letter(len(head) + 2)}2"
        if r.kind == "cycle" and len(head) >= 2:
            xl.line(ws, "容量 vs 循环", anchor, 1, [2], y_title="比容量")
            xl.heat(ws, 2, better="high")
        elif r.kind == "rate" and len(head) >= 2:
            xl.bar(ws, "倍率容量", anchor, 1, [2], y_title="比容量")
        elif r.kind == "eis" and len(head) >= 3:
            xl.line(ws, "Nyquist（Z' vs -Z''）", anchor, 2, [3], y_title="-Z''")
    return xl.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description="电化学数据 → docx/pptx/xlsx 三件套")
    ap.add_argument("--title", required=True)
    ap.add_argument("--in", dest="indir", required=True, help="数据目录（csv/xlsx）或单个文件")
    ap.add_argument("--out", required=True)
    ap.add_argument("--mass", type=float, default=None, help="活性物质质量 mg（把绝对容量换成 mAh/g）")
    ap.add_argument("--area", type=float, default=None, help="极片面积 cm²")
    ap.add_argument("--author", default="小蓝")
    ap.add_argument("--no-deck", action="store_true")
    a = ap.parse_args()

    os.makedirs(a.out, exist_ok=True)
    if os.path.isdir(a.indir):
        files = sorted(glob.glob(os.path.join(a.indir, "*.csv")) +
                       glob.glob(os.path.join(a.indir, "*.xlsx")))
    else:
        files = [a.indir]
    if not files:
        print(f"目录里没有 csv/xlsx：{a.indir}")
        return 2

    fig_dir = os.path.join(a.out, "figs")
    os.makedirs(fig_dir, exist_ok=True)
    items = []
    for f in files:
        print(f"· {os.path.basename(f)}")
        try:
            r, df = ec_data.analyze(f, mass_mg=a.mass, area_cm2=a.area)
        except Exception as e:  # noqa: BLE001
            print(f"  ! 读不了：{e}")
            continue
        r.source = f
        png = figure_for(r, df, os.path.join(fig_dir, os.path.splitext(os.path.basename(f))[0] + ".png"))
        print(f"  → {r.kind}：{headline(r)}{'' if png else '（无图）'}")
        items.append((r, df, png))
    if not items:
        print("没有可分析的数据")
        return 3
    items.sort(key=lambda t: KIND_ORDER.get(t[0].kind, 9))

    print("写出：", build_docx(items, a.title, os.path.join(a.out, "报告.docx"), a.author))
    print("写出：", build_xlsx(items, os.path.join(a.out, "数据底稿.xlsx")))
    if not a.no_deck:
        print("写出：", build_deck(items, a.title, os.path.join(a.out, "汇报.pptx"), a.author))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
